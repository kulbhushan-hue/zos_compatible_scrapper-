"""Excel report for vendor scrapes - same layout and IBM Carbon styling as the Broadcom example workbook.

Sheets: PTF & USERMOD Summary | Detail by release | Hyperlinks scraped | Pages searched | Notes
(An optional 'PTF Details' sheet can be added first with scrape_excel.add_ptf_sheet.)
"""

from __future__ import annotations

import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

import vendor_scraper as vs
from scrape_excel import (ACCENT, FIX_FILL, FONT, ZOS_ORDER, _body_style, _finish, _header, _link, _status,
                          add_ptf_sheet)


def _ids(r: vs.RowResult, rel: str, kind: str) -> str:
    ids = r.releases.get(rel)
    return ", ".join(getattr(ids, kind)) if ids else ""


def build_vendor_workbook(scrapes: list[vs.VendorScrape], ptf_table=None, title_note: str = "") -> bytes:
    wb = Workbook()

    # ------------------------------------------------------------- 1. summary
    ws = wb.active
    ws.title = "PTF & USERMOD Summary"
    head = ["Vendor", "Product", "VRM"]
    for z in ZOS_ORDER:
        head += [f"z/OS {z} status", f"z/OS {z} PTFs", f"z/OS {z} USERMOD / SYSMOD"]
    head += ["Release not stated: PTFs", "Release not stated: USERMOD / SYSMOD", "Note", "Source page"]
    ws.append(head)
    band = {4 + k * 3 + off: ("393939" if k % 2 == 0 else "262626") for k in range(3) for off in range(3)}
    _header(ws, band)
    for sc in scrapes:
        for r in sc.rows:
            row = [sc.vendor, r.product, r.vrm]
            for z in ZOS_ORDER:
                row += [r.status.get(z, ""), _ids(r, z, "ptfs"), _ids(r, z, "usermods")]
            row += [_ids(r, vs.UNSPECIFIED, "ptfs"), _ids(r, vs.UNSPECIFIED, "usermods"), r.note, r.link]
            ws.append(row)
    _body_style(ws)
    n = len(ZOS_ORDER)
    for i in range(2, ws.max_row + 1):
        for k in range(n):
            base = 4 + k * 3
            _status(ws.cell(row=i, column=base))
            for off in (1, 2):
                c = ws.cell(row=i, column=base + off)
                if c.value:
                    c.fill = PatternFill("solid", fgColor=FIX_FILL)
        for col in (4 + n * 3, 5 + n * 3):
            c = ws.cell(row=i, column=col)
            if c.value:
                c.fill = PatternFill("solid", fgColor=FIX_FILL)
        _link(ws.cell(row=i, column=len(head)), str(ws.cell(row=i, column=len(head)).value or ""))
    _finish(ws, [12, 40, 12] + [11, 28, 22] * 3 + [26, 26, 34, 40], "D2")

    # ------------------------------------------------------------- 2. detail
    ws = wb.create_sheet("Detail by release")
    cols = ["Vendor", "Product", "VRM", "z/OS release", "Status (dashboard data)", "PTFs", "USERMOD / SYSMOD IDs", "APAR IDs",
            "Additional PTFs (linked pages)", "Passage", "Section source", "Hyperlinks scraped", "Source page", "Note"]
    ws.append(cols)
    _header(ws)
    for sc in scrapes:
        for d in vs.flatten(sc):
            ws.append([d[c] for c in cols])
    _body_style(ws)
    for i in range(2, ws.max_row + 1):
        _status(ws.cell(row=i, column=5))
        for col in (6, 7, 8, 9):
            c = ws.cell(row=i, column=col)
            if c.value:
                c.fill = PatternFill("solid", fgColor=FIX_FILL)
        for col in (11, 12):
            c = ws.cell(row=i, column=col)
            first = str(c.value or "").split("\n")[0].split(" (page")[0]
            if first.startswith("http"):
                c.hyperlink = first
                c.font = Font(name=FONT, size=10, color=ACCENT, underline="single")
        _link(ws.cell(row=i, column=13), str(ws.cell(row=i, column=13).value or ""))
    _finish(ws, [11, 36, 12, 16, 14, 28, 22, 18, 28, 70, 44, 44, 40, 34], "E2")

    # ------------------------------------------------------------- 3. hyperlinks
    ws = wb.create_sheet("Hyperlinks scraped")
    ws.append(["Vendor", "URL", "Fetch status", "Page title", "PTFs found", "USERMOD / SYSMOD found", "APARs found", "Used by"])
    _header(ws)
    for sc in scrapes:
        users: dict[str, list[str]] = {}
        for r in sc.rows:
            for u in r.linked_urls:
                users.setdefault(u, []).append(f"{r.product} {r.vrm}")
        for url, p in sorted(sc.linked_pages.items()):
            who = users.get(url, [])
            ws.append([sc.vendor, url, p["status"], p["title"], ", ".join(p["ptfs"]), ", ".join(p["usermods"]), ", ".join(p["apars"]),
                       "; ".join(who[:6]) + (f" … (+{len(who) - 6} more)" if len(who) > 6 else "")])
    _body_style(ws)
    for i in range(2, ws.max_row + 1):
        _link(ws.cell(row=i, column=2), str(ws.cell(row=i, column=2).value or ""))
        for col in (5, 6, 7):
            if ws.cell(row=i, column=col).value:
                ws.cell(row=i, column=col).fill = PatternFill("solid", fgColor=FIX_FILL)
    _finish(ws, [12, 60, 30, 36, 28, 24, 22, 56], "C2")

    # ------------------------------------------------------------- 4. pages searched
    ws = wb.create_sheet("Pages searched")
    ws.append(["Vendor", "Source page", "Page no.", "URL searched", "Kind", "Status", "Passages read", "Note"])
    _header(ws)
    for sc in scrapes:
        for p in sc.pages:
            ws.append([sc.vendor, p.source_url, p.page_no, p.url, p.kind, p.status, p.blocks, p.note])
    _body_style(ws)
    red = PatternFill("solid", fgColor="FFD7D9")
    for i in range(2, ws.max_row + 1):
        _link(ws.cell(row=i, column=4), str(ws.cell(row=i, column=4).value or ""))
        if str(ws.cell(row=i, column=6).value).startswith("Login required"):
            ws.cell(row=i, column=6).fill = red
    _finish(ws, [12, 52, 9, 60, 7, 24, 12, 48], "C2")

    # ------------------------------------------------------------- 5. notes
    ws = wb.create_sheet("Notes")
    stamp = ", ".join(f"{sc.vendor} {sc.scraped_at}" for sc in scrapes)
    needs = [(sc.vendor, u) for sc in scrapes for u in sc.login_required]
    lines = [
        "Vendor PTF and USERMOD / SYSMOD extraction",
        f"Scraped at: {stamp}",
        f"Vendors: {', '.join(sc.vendor for sc in scrapes)}   |   products: {sum(len(sc.rows) for sc in scrapes)}   |   "
        f"pages searched: {sum(sum(1 for p in sc.pages if p.status == 'Scraped') for sc in scrapes)}",
    ]
    if title_note:
        lines.append(title_note)
    lines += [
        "",
        "How the values are produced",
        "1. Each vendor source page from the compatibility data is opened. If it has more than one page (next links, numbered pagination, ?page=N, or a multi-page PDF) at least 10 pages are searched when that many exist.",
        "2. Passages that name the product (and, when stated, its version) are read. PTF, USERMOD / SYSMOD and APAR IDs are taken only from passages that talk about fixes.",
        "3. IDs are assigned to z/OS 3.2 / 3.1 / 2.5 only when the passage names that release. Otherwise they are listed as 'Release not stated' and never attributed to a release.",
        "4. Hyperlinks inside matching passages are followed one level; IDs found there are listed as 'Additional PTFs (linked pages)' and must be verified.",
        "",
        "How to read it",
        "- Amber cells contain IDs. A blank cell means none was found in the text that could be read; it is NOT proof that no maintenance is required.",
        "- 'Login required' on the Pages searched sheet means the page could not be read without signing in. Supply the vendor user name and password in the dashboard and run again.",
        "- Pages rendered by JavaScript or protected by SSO / multi-factor authentication may return no readable text even after login; use a signed-in browser session cookie in that case.",
        "- Always confirm on the live vendor page and in your change process before applying maintenance.",
    ]
    if needs:
        lines += ["", "Sources that required a login in this run (credentials are never stored in this file):"] + [f"- {v}: {u}" for v, u in needs]
    for line in lines:
        ws.append([line])
    ws.column_dimensions["A"].width = 150
    for row in ws.iter_rows():
        for c in row:
            c.font = Font(name=FONT, size=10)
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws["A1"].font = Font(name=FONT, size=13, bold=True)
    for i in range(1, ws.max_row + 1):
        if ws.cell(row=i, column=1).value in ("How the values are produced", "How to read it"):
            ws.cell(row=i, column=1).font = Font(name=FONT, size=10, bold=True, color=ACCENT)

    if ptf_table is not None and len(ptf_table):
        add_ptf_sheet(wb, ptf_table, first=True)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
