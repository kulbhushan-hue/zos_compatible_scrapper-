"""Excel report for a Broadcom grid scrape (PTF + USERMOD/SYSMOD IDs per product and z/OS release).

Sheets
  1. PTF & USERMOD Summary  - one row per product/version, three z/OS releases side by side
  2. Detail by release      - one row per product + z/OS release, with the section text and source links
  3. Hyperlinks scraped     - every extra page that was followed, with the IDs found on it
  4. Notes                  - when/where it was scraped and how to read the values

Styling follows the IBM Carbon palette used by the dashboard (gray-100 header, blue-60 links,
green/red/yellow status tints). Hyperlinks are real, clickable Excel links.
"""

from __future__ import annotations

import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

import broadcom_scraper as bs

FONT = "Arial"
HEADER_FILL = "161616"
ACCENT = "0F62FE"
FIX_FILL = "FCF4D6"
STATUS_FILL = {"Yes": "DEFBE6", "No": "FFD7D9", "In Progress": "FCF4D6", "N/A": "F4F4F4", "Unverified": "E0E0E0"}
STATUS_FONT = {"Yes": "0E6027", "No": "A2191F", "In Progress": "684E00", "N/A": "525252", "Unverified": "525252"}
THIN = Side(style="thin", color="E0E0E0")
ZOS_ORDER = ["3.2", "3.1", "2.5"]


def _header(ws, fills: dict[int, str] | None = None):
    for i, c in enumerate(ws[1], 1):
        c.font = Font(name=FONT, bold=True, color="FFFFFF", size=10)
        c.fill = PatternFill("solid", fgColor=(fills or {}).get(i, HEADER_FILL))
        c.alignment = Alignment(vertical="center", horizontal="left", wrap_text=True)
        c.border = Border(bottom=Side(style="medium", color=ACCENT))
    ws.row_dimensions[1].height = 34


def _body_style(ws, first_row: int = 2):
    for row in ws.iter_rows(min_row=first_row):
        for c in row:
            c.font = Font(name=FONT, size=10)
            c.alignment = Alignment(wrap_text=True, vertical="top")
            c.border = Border(bottom=THIN)


def _link(cell, url: str, label: str | None = None):
    if url and url.lower().startswith(("http://", "https://")):
        if label is not None:
            cell.value = label
        cell.hyperlink = url
        cell.font = Font(name=FONT, size=10, color=ACCENT, underline="single")


def _status(cell):
    value = str(cell.value or "")
    if value in STATUS_FILL:
        cell.fill = PatternFill("solid", fgColor=STATUS_FILL[value])
        cell.font = Font(name=FONT, size=10, bold=True, color=STATUS_FONT[value])


def _finish(ws, widths: list[int], freeze: str, landscape: bool = True):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = freeze
    if ws.max_row > 1:
        ws.auto_filter.ref = ws.dimensions
    ws.sheet_view.zoomScale = 90
    ws.page_setup.orientation = "landscape" if landscape else "portrait"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = "1:1"


PTF_STATUS_FILL = {"Verified": "DEFBE6", "Partial": "FCF4D6", "Not found in source": "E0E0E0",
                   "Access denied": "FFD7D9", "Fetch failed": "FFD7D9", "Not attempted": "F4F4F4"}
PTF_STATUS_FONT = {"Verified": "0E6027", "Partial": "684E00", "Not found in source": "525252",
                   "Access denied": "A2191F", "Fetch failed": "A2191F", "Not attempted": "525252"}


def add_ptf_sheet(wb: Workbook, table, first: bool = False, note: str = ""):
    """'PTF Details' sheet: #, Product / Release, PTF ID, FMID, publish date, APAR, prerequisites,
    superseding PTF, full description, extraction status (+ catalog context columns)."""
    ws = wb.create_sheet("PTF Details", 0 if first else None)
    cols = list(table.columns)
    ws.append(cols)
    _header(ws)
    for row in table.itertuples(index=False):
        ws.append([("" if v is None else v) for v in row])
    _body_style(ws)
    s_idx = cols.index("Detail Extraction Status") + 1
    for r in range(2, ws.max_row + 1):
        c = ws.cell(row=r, column=s_idx)
        key = next((k for k in PTF_STATUS_FILL if str(c.value).startswith(k)), None)
        if key:
            c.fill = PatternFill("solid", fgColor=PTF_STATUS_FILL[key])
            c.font = Font(name=FONT, size=10, bold=True, color=PTF_STATUS_FONT[key])
        for name in ("FMID", "APAR ID", "Prerequisites", "Superseding PTF", "Full description", "PTF publish date"):
            cell = ws.cell(row=r, column=cols.index(name) + 1)
            if str(cell.value) in ("Not accessible", "Not found in source"):
                cell.font = Font(name=FONT, size=10, italic=True, color="6F6F6F")
        ws.cell(row=r, column=cols.index("PTF ID") + 1).font = Font(name=FONT, size=10, bold=True)
    widths = {"#": 5, "Product / Release": 44, "PTF ID": 11, "FMID": 14, "PTF publish date": 14, "APAR ID": 18,
              "Prerequisites": 24, "Superseding PTF": 18, "Full description": 50, "Detail Extraction Status": 34,
              "z/OS release (matrix)": 12, "Catalog description": 50, "Feature": 8, "HIPER": 7, "PE": 6,
              "CARSID": 10, "SECINT": 8, "Detail source": 40}
    _finish(ws, [widths.get(c, 14) for c in cols], "D2")
    if note:
        ws.cell(row=ws.max_row + 2, column=2, value=note).font = Font(name=FONT, size=9, italic=True, color="525252")
    return ws


def build_ptf_workbook(table, note: str = "") -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    add_ptf_sheet(wb, table, note=note or "Not accessible = the detail page could not be read with the access available. Not found in source = the source was read but does not contain the field. Neither means 'none'.")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_scrape_workbook(result: bs.ScrapeResult, title_note: str = "", ptf_table=None) -> bytes:
    wb = Workbook()

    # ------------------------------------------------------------- 1. summary
    ws = wb.active
    ws.title = "PTF & USERMOD Summary"
    head = ["Vendor", "Product", "VRM"]
    for z in ZOS_ORDER:
        head += [f"z/OS {z} status", f"z/OS {z} PTFs", f"z/OS {z} USERMOD / SYSMOD"]
    head += ["Product page"]
    ws.append(head)
    band = {}
    for k, z in enumerate(ZOS_ORDER):                       # alternate header tint per release block
        for off in range(3):
            band[4 + k * 3 + off] = "393939" if k % 2 == 0 else "262626"
    _header(ws, band)
    for rec in result.records:
        product, vrm = bs.split_name_vrm(rec.clean_name)
        row = ["Broadcom", product + (" (End of service/life)" if rec.name != rec.clean_name else ""), vrm]
        for z in ZOS_ORDER:
            res = rec.releases.get(z)
            row += [res.status if res else "", ", ".join(res.ptfs) if res else "", ", ".join(res.usermods) if res else ""]
        row.append(rec.product_url)
        ws.append(row)
    _body_style(ws)
    for r in range(2, ws.max_row + 1):
        for k in range(len(ZOS_ORDER)):
            base = 4 + k * 3
            _status(ws.cell(row=r, column=base))
            for off in (1, 2):
                c = ws.cell(row=r, column=base + off)
                if c.value:
                    c.fill = PatternFill("solid", fgColor=FIX_FILL)
        _link(ws.cell(row=r, column=len(head)), str(ws.cell(row=r, column=len(head)).value or ""), "Open product page")
    _finish(ws, [12, 44, 10] + [11, 30, 22] * 3 + [18], "D2")

    # ------------------------------------------------------------- 2. detail
    ws = wb.create_sheet("Detail by release")
    rows = bs.flatten(result)
    cols = ["Vendor", "Product", "VRM", "z/OS release", "Status (grid)", "PTFs (grid sections)", "USERMOD / SYSMOD IDs",
            "Additional PTFs (linked pages)", "IBM z/OS service IDs mentioned", "Grid links", "Section notes",
            "Section source", "Hyperlinks scraped", "Product page"]
    ws.append(cols)
    _header(ws)
    for r in rows:
        ws.append([r[c] for c in cols])
    _body_style(ws)
    for i in range(2, ws.max_row + 1):
        _status(ws.cell(row=i, column=5))
        for col in (6, 7, 8):
            c = ws.cell(row=i, column=col)
            if c.value:
                c.fill = PatternFill("solid", fgColor=FIX_FILL)
        for col in (12, 13):                                 # newline-separated URL lists: link the first URL
            c = ws.cell(row=i, column=col)
            first = str(c.value or "").split("\n")[0]
            if first.startswith("http"):
                c.hyperlink = first
                c.font = Font(name=FONT, size=10, color=ACCENT, underline="single")
        _link(ws.cell(row=i, column=14), str(ws.cell(row=i, column=14).value or ""))
    _finish(ws, [10, 36, 9, 10, 12, 34, 24, 30, 20, 22, 70, 46, 46, 40], "E2")

    # ------------------------------------------------------------- 3. hyperlinks
    ws = wb.create_sheet("Hyperlinks scraped")
    ws.append(["URL", "Fetch status", "Page title", "PTFs found", "USERMOD / SYSMOD found", "IBM z/OS service IDs", "Used by"])
    _header(ws)
    users: dict[str, list[str]] = {}
    for rec in result.records:
        for z, res in rec.releases.items():
            for u in res.linked_urls:
                users.setdefault(u, []).append(f"{rec.clean_name} (z/OS {z})")
        if rec.product_url:
            users.setdefault(rec.product_url, []).append(f"{rec.clean_name} (product page)")
    for url, page in sorted(result.linked_pages.items()):
        who = users.get(url, [])
        ws.append([url, page.status, page.title, ", ".join(page.ptfs), ", ".join(page.usermods), ", ".join(page.other_ids),
                   "; ".join(who[:6]) + (f" … (+{len(who) - 6} more)" if len(who) > 6 else "")])
    _body_style(ws)
    for i in range(2, ws.max_row + 1):
        _link(ws.cell(row=i, column=1), str(ws.cell(row=i, column=1).value or ""))
        for col in (4, 5):
            if ws.cell(row=i, column=col).value:
                ws.cell(row=i, column=col).fill = PatternFill("solid", fgColor=FIX_FILL)
    _finish(ws, [60, 28, 36, 30, 24, 22, 60], "B2")

    # ------------------------------------------------------------- 4. notes
    ws = wb.create_sheet("Notes")
    lines = [
        "Broadcom z/OS compatibility grid - PTF and USERMOD / SYSMOD extraction",
        f"Source page: {result.grid_url}",
        f"Scraped at: {result.scraped_at}   |   products on the grid page: {len(result.records)}   |   extra pages followed: {len(result.linked_pages)}",
    ]
    if title_note:
        lines.append(title_note)
    lines += [
        "",
        "How the values are produced",
        "1. The grid table is read row by row: product + version, then the z/OS V3R2, V3R1 and V2R5 status cells.",
        "2. The 'PTFs' and 'Additional Information' links in each cell point to a section lower on the same page. That section's text is read and the IDs in it are extracted.",
        "3. Hyperlinks inside those sections (and the product page named in the first column) are fetched one level deep; IDs found there are listed separately as 'Additional PTFs (linked pages)'.",
        "4. Broadcom PTF / solution IDs have the pattern two letters + five digits (for example LU18187, RO94173). IDs on a line that names a USERMOD or SYSMOD are reported as USERMOD / SYSMOD IDs.",
        "5. IBM z/OS service IDs (UA..., OA..., PH...) that appear in Broadcom text are z/OS maintenance, not Broadcom PTFs; they are listed in their own column.",
        "",
        "How to read it",
        "- Amber cells contain IDs. A blank ID cell means none was found in the fetched text; it is NOT proof that no maintenance is required.",
        "- 'Additional PTFs (linked pages)' may belong to another release of the same product. Verify before applying.",
        "- Linked pages behind a login or rendered by JavaScript (for example parts of support.broadcom.com) can return no readable text; the status column on the 'Hyperlinks scraped' sheet says so.",
        "- Always confirm on the live Broadcom page and in your change process before applying maintenance.",
    ]
    for line in lines:
        ws.append([line])
    ws.column_dimensions["A"].width = 150
    for row in ws.iter_rows():
        for c in row:
            c.font = Font(name=FONT, size=10)
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws["A1"].font = Font(name=FONT, size=13, bold=True)
    for i in range(1, ws.max_row + 1):
        if ws.cell(row=i, column=1).value in ("How the values are produced", "How to read it"):
            ws.cell(row=i, column=1).font = Font(name=FONT, size=10, bold=True, color=ACCENT)
    _link(ws["A2"], result.grid_url, f"Source page: {result.grid_url}")

    if ptf_table is not None and len(ptf_table):
        add_ptf_sheet(wb, ptf_table, first=True)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
